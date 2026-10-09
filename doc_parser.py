import io
import os
import zipfile
from typing import Dict, Any, Tuple

import pypdf
import docx
import pptx

def extract_text_from_pdf(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """Extract text from all pages of a PDF document."""
    reader = pypdf.PdfReader(io.BytesIO(file_bytes))
    num_pages = len(reader.pages)
    text_pages = []
    
    for idx, page in enumerate(reader.pages):
        page_text = page.extract_text() or ""
        if page_text.strip():
            text_pages.append(page_text.strip())
            
    full_text = "\n\n".join(text_pages)
    meta = {
        "format": "pdf",
        "page_count": num_pages,
        "extracted_pages": len(text_pages)
    }
    return full_text, meta

def extract_text_from_docx(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """Extract text from Word .docx document."""
    doc = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    
    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)
                
    full_text = "\n\n".join(paragraphs)
    meta = {
        "format": "docx",
        "paragraph_count": len(paragraphs)
    }
    return full_text, meta

def extract_text_from_pptx(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """Extract text from PowerPoint .pptx presentation slides."""
    prs = pptx.Presentation(io.BytesIO(file_bytes))
    slide_texts = []
    
    for idx, slide in enumerate(prs.slides):
        slide_lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        slide_lines.append(line)
        if slide_lines:
            slide_texts.append(f"--- Slide {idx+1} ---\n" + "\n".join(slide_lines))
            
    full_text = "\n\n".join(slide_texts)
    meta = {
        "format": "pptx",
        "slide_count": len(prs.slides),
        "extracted_slides": len(slide_texts)
    }
    return full_text, meta

def extract_text_from_zip(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """
    Decompress .zip archive in memory and extract text from all contained PDF, DOCX, PPTX, TXT files.
    """
    extracted_sections = []
    processed_files = []
    total_pages = 0

    with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
        file_idx = 0
        for zip_info in zf.infolist():
            if zip_info.is_dir():
                continue
            fname = zip_info.filename
            if fname.startswith("__MACOSX") or "/." in fname or fname.startswith("."):
                continue

            ext = os.path.splitext(fname)[1].lower()
            if ext in [".pdf", ".docx", ".doc", ".pptx", ".ppt", ".txt", ".md", ".csv", ".json"]:
                with zf.open(zip_info) as f:
                    child_bytes = f.read()
                try:
                    child_text, child_meta = parse_uploaded_file(fname, child_bytes)
                    if child_text.strip():
                        file_idx += 1
                        entry_id = f"file_{file_idx}"
                        words_cnt = len(child_text.split())
                        pages_cnt = child_meta.get("page_count", child_meta.get("slide_count", 1))
                        
                        extracted_sections.append(f"=== Document: {fname} ===\n" + child_text.strip())
                        processed_files.append({
                            "id": entry_id,
                            "filename": fname,
                            "format": child_meta.get("format", ext.lstrip('.')),
                            "words": words_cnt,
                            "pages": pages_cnt,
                            "text": child_text.strip(),
                            "meta": child_meta
                        })
                        total_pages += pages_cnt
                except Exception as e:
                    print(f"Skipping corrupt file in zip {fname}: {e}")

    full_text = "\n\n".join(extracted_sections)
    meta = {
        "format": "zip",
        "zip_file_count": len(processed_files),
        "processed_files": processed_files,
        "page_count": total_pages or len(processed_files)
    }
    return full_text, meta

def parse_uploaded_file(filename: str, file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """Determine file type by extension and extract text."""
    ext = os.path.splitext(filename)[1].lower()
    
    if ext == ".zip":
        return extract_text_from_zip(file_bytes)
    elif ext == ".pdf":
        return extract_text_from_pdf(file_bytes)
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(file_bytes)
    elif ext in [".pptx", ".ppt"]:
        return extract_text_from_pptx(file_bytes)
    elif ext in [".txt", ".md", ".csv", ".json", ".html"]:
        text = file_bytes.decode("utf-8", errors="ignore")
        return text, {"format": ext.lstrip('.')}
    else:
        # Fallback to UTF-8 text decoding
        text = file_bytes.decode("utf-8", errors="ignore")
        return text, {"format": "unknown"}
