import os
from io import BytesIO
from typing import Tuple
from pypdf import PdfReader
import docx
import pptx

def parse_document(file_bytes: bytes, filename: str) -> Tuple[str, dict]:
    """
    Extract text and metadata from PDF, DOCX, PPTX, TXT files.
    Returns (extracted_text, metadata_dict).
    """
    ext = os.path.splitext(filename)[1].lower()
    metadata = {"filename": filename, "format": ext, "pages": 1}
    
    if ext == ".pdf":
        reader = PdfReader(BytesIO(file_bytes))
        metadata["pages"] = len(reader.pages)
        pages_text = []
        for idx, page in enumerate(reader.pages):
            txt = page.extract_text()
            if txt and txt.strip():
                pages_text.append(txt.strip())
        full_text = "\n\n".join(pages_text)
        
    elif ext in [".docx", ".doc"]:
        doc = docx.Document(BytesIO(file_bytes))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        metadata["paragraphs"] = len(paragraphs)
        full_text = "\n\n".join(paragraphs)
        
    elif ext in [".pptx", ".ppt"]:
        prs = pptx.Presentation(BytesIO(file_bytes))
        metadata["pages"] = len(prs.slides) # slides count
        slide_texts = []
        for idx, slide in enumerate(prs.slides):
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
            if texts:
                slide_texts.append(f"[Slide {idx + 1}]\n" + "\n".join(texts))
        full_text = "\n\n".join(slide_texts)
        
    else: # Default text / markdown / txt
        try:
            full_text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            full_text = file_bytes.decode("latin-1", errors="ignore")
            
    return full_text, metadata
